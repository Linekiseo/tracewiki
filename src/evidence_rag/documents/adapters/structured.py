from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass, field
from enum import StrEnum
from html.parser import HTMLParser
from pathlib import Path
from statistics import median
from typing import Any, Protocol
from urllib.parse import unquote

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf import __version__ as pypdf_version
from pypdf.errors import PdfReadError

_SECRET_RE = re.compile(r"(?i)(?:api[_-]?key|access[_-]?token|secret|password)\s*[:=]\s*[^\s,;]+")
_ABS_PATH_RE = re.compile(
    r"(?i)(?:^|[\s\"'`(])(?:/(?:Users|home|private|tmp|var|etc|opt)/"
    r"|[a-z]:[\\/]|\\\\[^\\/\s]+[\\/][^\\/\s]+)"
)
_ENCODED_PATH_MARKERS = ("%2fusers%2f", "%2fhome%2f", "%5c%5c", "%2f%2f")
_PROVIDER_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")
_SCAN_TEXT_THRESHOLD = 24


class PDFProviderStatus(StrEnum):
    AVAILABLE = "available"
    PARTIAL = "partial"
    UNAVAILABLE = "unavailable"
    ERROR = "error"


class PDFOcrStatus(StrEnum):
    NOT_REQUIRED = "not_required"
    AVAILABLE = "available"
    PARTIAL = "partial"
    UNAVAILABLE = "unavailable"
    ERROR = "error"


class DocumentSecurityError(ValueError):
    """Raised when a document contains active or protected content."""


@dataclass(frozen=True, slots=True)
class PDFLayoutBlock:
    text: str
    bbox: tuple[float, float, float, float]
    block_type: str = "body"
    confidence: float = 1.0


@dataclass(frozen=True, slots=True)
class PDFPageRequest:
    page_number: int
    page_identity: str
    width: float
    height: float
    native_text: str


@dataclass(frozen=True, slots=True)
class PDFParseRequest:
    artifact_sha256: str
    pdf_bytes: bytes = field(repr=False)
    pages: tuple[PDFPageRequest, ...] = ()


@dataclass(frozen=True, slots=True)
class PDFProviderPage:
    page_number: int
    page_identity: str
    status: PDFProviderStatus
    ocr_text: str = ""
    blocks: tuple[PDFLayoutBlock, ...] = ()


@dataclass(frozen=True, slots=True)
class PDFProviderResult:
    provider_name: str
    provider_version: str
    status: PDFProviderStatus
    pages: tuple[PDFProviderPage, ...] = ()


class PDFParserProvider(Protocol):
    provider_name: str
    provider_version: str

    def parse(self, request: PDFParseRequest) -> PDFProviderResult: ...


class UnavailablePDFParserProvider:
    """Default local provider: explicitly reports that OCR is not installed."""

    provider_name = "local-unavailable"
    provider_version = "unavailable-v1"

    def parse(self, request: PDFParseRequest) -> PDFProviderResult:
        return PDFProviderResult(
            provider_name=self.provider_name,
            provider_version=self.provider_version,
            status=PDFProviderStatus.UNAVAILABLE,
        )


@dataclass(frozen=True, slots=True)
class PDFDeterministicPage:
    """Injectable local output used by deterministic/offline providers and tests."""

    ocr_text: str
    blocks: tuple[PDFLayoutBlock, ...]
    status: PDFProviderStatus = PDFProviderStatus.AVAILABLE


class DeterministicPDFParserProvider:
    """Offline provider whose page outputs are supplied explicitly by the caller."""

    provider_name = "deterministic-local"
    provider_version = "fixture-v1"

    def __init__(
        self,
        pages: dict[int, PDFDeterministicPage],
        *,
        result_status: PDFProviderStatus | None = None,
    ) -> None:
        self._pages = dict(pages)
        self._result_status = result_status

    def parse(self, request: PDFParseRequest) -> PDFProviderResult:
        requested = {item.page_number: item for item in request.pages}
        pages = tuple(
            PDFProviderPage(
                page_number=page_number,
                page_identity=requested[page_number].page_identity,
                status=output.status,
                ocr_text=output.ocr_text,
                blocks=output.blocks,
            )
            for page_number, output in sorted(self._pages.items())
            if page_number in requested
        )
        if self._result_status is not None:
            status = self._result_status
        elif not pages:
            status = PDFProviderStatus.UNAVAILABLE
        elif len(pages) < len(requested) or any(
            item.status != PDFProviderStatus.AVAILABLE for item in pages
        ):
            status = PDFProviderStatus.PARTIAL
        else:
            status = PDFProviderStatus.AVAILABLE
        return PDFProviderResult(
            provider_name=self.provider_name,
            provider_version=self.provider_version,
            status=status,
            pages=pages,
        )


def _canonical_sha256(value: object) -> str:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        allow_nan=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    return "sha256:" + hashlib.sha256(payload).hexdigest()


def _bytes_sha256(value: bytes) -> str:
    return "sha256:" + hashlib.sha256(value).hexdigest()


def _text_sha256(value: str) -> str:
    return _bytes_sha256(value.encode("utf-8"))


def _safe_provider_id(value: object) -> str:
    candidate = str(value)
    if _PROVIDER_ID_RE.fullmatch(candidate):
        return candidate
    return "invalid-provider"


def _reject_protected_text(value: str) -> None:
    decoded = value
    for _ in range(4):
        candidate = unquote(decoded)
        if candidate == decoded:
            break
        decoded = candidate
    lowered = decoded.casefold()
    if (
        _SECRET_RE.search(decoded)
        or _ABS_PATH_RE.search(decoded)
        or any(marker in lowered for marker in _ENCODED_PATH_MARKERS)
    ):
        raise DocumentSecurityError("document contains protected secret or absolute path")


def _validate_bbox(
    bbox: tuple[float, float, float, float],
    width: float,
    height: float,
) -> None:
    if len(bbox) != 4 or not all(math.isfinite(float(value)) for value in bbox):
        raise ValueError("parser block bbox is invalid")
    x0, top, x1, bottom = (float(value) for value in bbox)
    tolerance = 1e-6
    if (
        x0 < -tolerance
        or top < -tolerance
        or x1 > width + tolerance
        or bottom > height + tolerance
        or x1 <= x0
        or bottom <= top
    ):
        raise ValueError("parser block bbox is outside its page")


def _estimated_text_width(text: str, font_size: float) -> float:
    factors = []
    for character in text:
        if character.isspace() or character in "ilI.,:;'|!":
            factors.append(0.28)
        elif character in "MW@#%&":
            factors.append(0.88)
        elif character.isupper():
            factors.append(0.62)
        else:
            factors.append(0.52)
    return max(font_size * 0.25, font_size * sum(factors))


def _pdf_object(value: Any) -> Any:
    try:
        return value.get_object()
    except AttributeError:
        return value


@dataclass(slots=True)
class ParsedDocument:
    content: str
    source_type: str
    source_uri: str
    sections: list[dict[str, Any]] = field(default_factory=list)
    pages: list[dict[str, Any]] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    figures: list[dict[str, Any]] = field(default_factory=list)
    citations: list[dict[str, Any]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


class StructuredDocumentAdapter:
    adapter_version = "structured-document-v2"

    def __init__(self, pdf_provider: PDFParserProvider | None = None) -> None:
        self.pdf_provider = pdf_provider or UnavailablePDFParserProvider()

    def parse(self, path: Path) -> ParsedDocument:
        suffix = path.suffix.casefold()
        if suffix in {".md", ".markdown"}:
            return self._markdown(path)
        if suffix == ".txt":
            return self._text(path)
        if suffix == ".pdf":
            return self._pdf(path)
        if suffix == ".docx":
            return self._docx(path)
        if suffix in {".html", ".htm"}:
            return self._html(path)
        raise ValueError(f"unsupported document format: {suffix or '<none>'}")

    def inline(self, content: str, source_uri: str) -> ParsedDocument:
        sections = self._line_sections(content, source_uri)
        return ParsedDocument(
            content=content,
            source_type="inline_text",
            source_uri=source_uri,
            sections=sections,
            tables=self._markdown_tables(content, source_uri),
            citations=self._citations(content, source_uri),
        )

    def _markdown(self, path: Path) -> ParsedDocument:
        content = path.read_text(encoding="utf-8")
        source_uri = str(path)
        return ParsedDocument(
            content=content,
            source_type="markdown",
            source_uri=source_uri,
            sections=self._line_sections(content, source_uri),
            tables=self._markdown_tables(content, source_uri),
            citations=self._citations(content, source_uri),
        )

    def _text(self, path: Path) -> ParsedDocument:
        content = path.read_text(encoding="utf-8")
        source_uri = str(path)
        return ParsedDocument(
            content=content,
            source_type="text",
            source_uri=source_uri,
            sections=self._line_sections(content, source_uri),
            citations=self._citations(content, source_uri),
        )

    def _pdf(self, path: Path) -> ParsedDocument:
        raw = path.read_bytes()
        artifact_sha256 = _bytes_sha256(raw)
        source_uri = f"document-source://sha256/{artifact_sha256.removeprefix('sha256:')}"
        try:
            reader = PdfReader(path)
        except PdfReadError as error:
            raise ValueError("document is not a valid PDF") from error
        self._reject_unsafe_pdf(reader)
        native_pages: list[dict[str, Any]] = []
        scan_requests: list[PDFPageRequest] = []
        for index, page in enumerate(reader.pages, start=1):
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            content_bytes = self._page_content_bytes(page)
            page_content_sha256 = _bytes_sha256(content_bytes)
            page_identity = _canonical_sha256(
                {
                    "artifact_sha256": artifact_sha256,
                    "page_number": index,
                    "page_content_sha256": page_content_sha256,
                    "width": width,
                    "height": height,
                }
            )
            native_text, native_blocks = self._native_layout(
                page,
                page_number=index,
                page_identity=page_identity,
                width=width,
                height=height,
            )
            _reject_protected_text(native_text)
            image_count = self._image_count(page)
            native_char_count = len(re.sub(r"\s+", "", native_text))
            scan_detected = image_count > 0 and native_char_count < _SCAN_TEXT_THRESHOLD
            scan_reason = (
                "image_page_with_insufficient_native_text"
                if scan_detected
                else "native_text_or_non_image_page"
            )
            record = {
                "page_number": index,
                "page_identity": page_identity,
                "page_content_sha256": page_content_sha256,
                "width": width,
                "height": height,
                "native_text": native_text,
                "native_blocks": native_blocks,
                "scan_detected": scan_detected,
                "scan_detection": {
                    "algorithm": "image-low-text-v1",
                    "image_count": image_count,
                    "native_character_count": native_char_count,
                    "threshold": _SCAN_TEXT_THRESHOLD,
                    "reason": scan_reason,
                },
            }
            native_pages.append(record)
            if scan_detected:
                scan_requests.append(
                    PDFPageRequest(
                        page_number=index,
                        page_identity=page_identity,
                        width=width,
                        height=height,
                        native_text=native_text,
                    )
                )

        provider_result, provider_input_sha256, provider_output_sha256 = self._provider_result(
            raw=raw,
            artifact_sha256=artifact_sha256,
            pages=tuple(scan_requests),
        )
        provider_pages = {item.page_number: item for item in provider_result.pages}
        pages: list[dict[str, Any]] = []
        sections: list[dict[str, Any]] = []
        content_parts: list[str] = []
        ocr_counts = {status.value: 0 for status in PDFOcrStatus}
        for record in native_pages:
            index = int(record["page_number"])
            scan_detected = bool(record["scan_detected"])
            provider_page = provider_pages.get(index)
            if not scan_detected:
                ocr_status = PDFOcrStatus.NOT_REQUIRED
                content = str(record["native_text"])
                selected_blocks = record["native_blocks"]
            elif provider_page is None:
                ocr_status = (
                    PDFOcrStatus.ERROR
                    if provider_result.status == PDFProviderStatus.ERROR
                    else PDFOcrStatus.UNAVAILABLE
                )
                content = str(record["native_text"])
                selected_blocks = record["native_blocks"]
            else:
                ocr_status = {
                    PDFProviderStatus.AVAILABLE: PDFOcrStatus.AVAILABLE,
                    PDFProviderStatus.PARTIAL: PDFOcrStatus.PARTIAL,
                    PDFProviderStatus.UNAVAILABLE: PDFOcrStatus.UNAVAILABLE,
                    PDFProviderStatus.ERROR: PDFOcrStatus.ERROR,
                }[provider_page.status]
                content = provider_page.ocr_text
                selected_blocks = provider_page.blocks
            ocr_counts[ocr_status.value] += 1
            locator = f"{source_uri}/page/{index}"
            serialized_blocks = self._serialize_blocks(
                selected_blocks,
                page_identity=str(record["page_identity"]),
                width=float(record["width"]),
                height=float(record["height"]),
                source="ocr" if scan_detected and provider_page is not None else "native",
                parser_name=(
                    provider_result.provider_name
                    if scan_detected and provider_page is not None
                    else "pypdf"
                ),
                parser_version=(
                    provider_result.provider_version
                    if scan_detected and provider_page is not None
                    else pypdf_version
                ),
                bbox_method=(
                    "provider-supplied"
                    if scan_detected and provider_page is not None
                    else "pdf-text-matrix-v1"
                ),
            )
            layout_payload_sha256 = _canonical_sha256(
                {
                    "page_identity": record["page_identity"],
                    "blocks": serialized_blocks,
                }
            )
            metadata: dict[str, Any] = {
                "width": record["width"],
                "height": record["height"],
                "artifact_sha256": artifact_sha256,
                "page_identity": record["page_identity"],
                "page_content_sha256": record["page_content_sha256"],
                "native_text_sha256": _text_sha256(str(record["native_text"])),
                "scan_detected": scan_detected,
                "scan_detection": record["scan_detection"],
                "ocr_status": ocr_status.value,
                "text_source": (
                    "ocr"
                    if scan_detected and provider_page is not None and provider_page.ocr_text
                    else "native"
                ),
                "blocks": serialized_blocks,
                "layout_payload_sha256": layout_payload_sha256,
                "native_layout": {
                    "parser_name": "pypdf",
                    "parser_version": pypdf_version,
                    "bbox_method": "pdf-text-matrix-v1",
                },
            }
            if scan_detected:
                metadata["parser_provenance"] = {
                    "provider_name": provider_result.provider_name,
                    "provider_version": provider_result.provider_version,
                    "provider_status": provider_result.status.value,
                    "input_sha256": provider_input_sha256,
                    "output_sha256": provider_output_sha256,
                }
            pages.append(
                {
                    "page_number": index,
                    "content": content,
                    "source_locator": locator,
                    "metadata": metadata,
                }
            )
            sections.append(
                {
                    "ordinal": index,
                    "level": 1,
                    "title": f"Page {index}",
                    "content": content,
                    "start_line": index,
                    "end_line": index,
                    "source_locator": locator,
                    "page_number": index,
                }
            )
            if content:
                content_parts.append(content)
        content = "\n\n".join(content_parts)
        return ParsedDocument(
            content=content,
            source_type="pdf",
            source_uri=source_uri,
            pages=pages,
            sections=sections,
            citations=self._citations(content, source_uri),
            metadata={
                "page_count": len(pages),
                "artifact_sha256": artifact_sha256,
                "layout_parser": "pypdf-layout-v1",
                "provider_status": provider_result.status.value,
                "ocr_status_counts": ocr_counts,
            },
        )

    def _provider_result(
        self,
        *,
        raw: bytes,
        artifact_sha256: str,
        pages: tuple[PDFPageRequest, ...],
    ) -> tuple[PDFProviderResult, str, str]:
        input_payload = {
            "artifact_sha256": artifact_sha256,
            "pages": [
                {
                    "page_number": item.page_number,
                    "page_identity": item.page_identity,
                    "width": item.width,
                    "height": item.height,
                    "native_text_sha256": _text_sha256(item.native_text),
                }
                for item in pages
            ],
        }
        input_sha256 = _canonical_sha256(input_payload)
        if not pages:
            result = PDFProviderResult(
                provider_name=self.pdf_provider.provider_name,
                provider_version=self.pdf_provider.provider_version,
                status=PDFProviderStatus.AVAILABLE,
            )
            payload = self._provider_payload(result)
            return result, input_sha256, _canonical_sha256(payload)
        request = PDFParseRequest(
            artifact_sha256=artifact_sha256,
            pdf_bytes=raw,
            pages=pages,
        )
        try:
            result = self.pdf_provider.parse(request)
            self._validate_provider_result(result, request)
        except DocumentSecurityError:
            raise
        except Exception:
            result = PDFProviderResult(
                provider_name=_safe_provider_id(
                    getattr(self.pdf_provider, "provider_name", "local-provider")
                ),
                provider_version=_safe_provider_id(
                    getattr(self.pdf_provider, "provider_version", "error-v1")
                ),
                status=PDFProviderStatus.ERROR,
            )
        payload = self._provider_payload(result)
        return result, input_sha256, _canonical_sha256(payload)

    @staticmethod
    def _provider_payload(result: PDFProviderResult) -> dict[str, Any]:
        return {
            "provider_name": result.provider_name,
            "provider_version": result.provider_version,
            "status": result.status.value,
            "pages": [
                {
                    "page_number": item.page_number,
                    "page_identity": item.page_identity,
                    "status": item.status.value,
                    "ocr_text_sha256": _text_sha256(item.ocr_text),
                    "blocks": [
                        {
                            "text_sha256": _text_sha256(block.text),
                            "bbox": list(block.bbox),
                            "block_type": block.block_type,
                            "confidence": block.confidence,
                        }
                        for block in item.blocks
                    ],
                }
                for item in result.pages
            ],
        }

    @staticmethod
    def _validate_provider_result(
        result: PDFProviderResult,
        request: PDFParseRequest,
    ) -> None:
        if not _PROVIDER_ID_RE.fullmatch(result.provider_name) or not _PROVIDER_ID_RE.fullmatch(
            result.provider_version
        ):
            raise ValueError("parser provider identity is invalid")
        expected = {item.page_number: item for item in request.pages}
        seen: set[int] = set()
        for page in result.pages:
            source = expected.get(page.page_number)
            if source is None or page.page_number in seen:
                raise ValueError("parser provider page membership is invalid")
            seen.add(page.page_number)
            if page.page_identity != source.page_identity:
                raise ValueError("parser provider page identity mismatch")
            if page.status == PDFProviderStatus.AVAILABLE and not page.ocr_text.strip():
                raise ValueError("available parser page has no OCR text")
            if page.status in {PDFProviderStatus.UNAVAILABLE, PDFProviderStatus.ERROR} and (
                page.ocr_text or page.blocks
            ):
                raise ValueError("unavailable parser page returned derived content")
            _reject_protected_text(page.ocr_text)
            for block in page.blocks:
                _reject_protected_text(block.text)
                _validate_bbox(block.bbox, source.width, source.height)
                if block.block_type not in {"header", "body", "footer", "table", "figure"}:
                    raise ValueError("parser provider block type is invalid")
                if not 0.0 <= block.confidence <= 1.0:
                    raise ValueError("parser provider confidence is invalid")
        if (
            result.status in {PDFProviderStatus.UNAVAILABLE, PDFProviderStatus.ERROR}
            and result.pages
        ):
            raise ValueError("unavailable parser provider returned page payloads")
        if result.status == PDFProviderStatus.AVAILABLE and seen != set(expected):
            raise ValueError("available parser provider omitted requested pages")

    @staticmethod
    def _serialize_blocks(
        blocks: tuple[PDFLayoutBlock, ...] | list[PDFLayoutBlock],
        *,
        page_identity: str,
        width: float,
        height: float,
        source: str,
        parser_name: str,
        parser_version: str,
        bbox_method: str,
    ) -> list[dict[str, Any]]:
        serialized: list[dict[str, Any]] = []
        for ordinal, block in enumerate(blocks, 1):
            _validate_bbox(block.bbox, width, height)
            text = block.text.strip()
            if not text:
                continue
            value = {
                "block_id": hashlib.sha256(
                    f"{page_identity}\x1f{ordinal}\x1f{text}\x1f{block.bbox}".encode()
                ).hexdigest()[:32],
                "ordinal": ordinal,
                "block_type": block.block_type,
                "text": text,
                "bbox": list(block.bbox),
                "text_sha256": _text_sha256(text),
                "source": source,
                "confidence": block.confidence,
                "parser_name": _safe_provider_id(parser_name),
                "parser_version": _safe_provider_id(parser_version),
                "bbox_method": _safe_provider_id(bbox_method),
            }
            serialized.append(value)
        return serialized

    @staticmethod
    def _native_layout(
        page: Any,
        *,
        page_number: int,
        page_identity: str,
        width: float,
        height: float,
    ) -> tuple[str, tuple[PDFLayoutBlock, ...]]:
        fragments: list[dict[str, Any]] = []

        def visitor(
            text: str,
            cm: list[float],
            tm: list[float],
            font: dict[str, Any] | None,
            font_size: float,
        ) -> None:
            value = " ".join(text.split())
            if not value:
                return
            x = float(tm[4] * cm[0] + tm[5] * cm[2] + cm[4])
            baseline = float(tm[4] * cm[1] + tm[5] * cm[3] + cm[5])
            size = max(1.0, abs(float(font_size)))
            font_name = str((font or {}).get("/BaseFont") or "")
            estimated_width = _estimated_text_width(value, size)
            x0 = min(max(0.0, x), width)
            x1 = min(width, max(x0 + 0.1, x + estimated_width))
            top = min(height, max(0.0, height - baseline - size))
            bottom = min(height, max(top + 0.1, top + size * 1.2))
            if x1 <= x0 or bottom <= top:
                return
            fragments.append(
                {
                    "text": value,
                    "bbox": (x0, top, x1, bottom),
                    "font_size": size,
                    "font_name": font_name,
                }
            )

        text = (page.extract_text(visitor_text=visitor) or "").strip()
        if not fragments:
            return text, ()
        body_size = median(item["font_size"] for item in fragments)
        ordered = sorted(fragments, key=lambda item: (item["bbox"][1], item["bbox"][0]))
        blocks = tuple(
            PDFLayoutBlock(
                text=item["text"],
                bbox=item["bbox"],
                block_type=(
                    "header"
                    if item["font_size"] >= body_size * 1.18
                    or "bold" in item["font_name"].casefold()
                    else "body"
                ),
                confidence=0.96,
            )
            for item in ordered
        )
        return text, blocks

    @staticmethod
    def _image_count(page: Any) -> int:
        resources = _pdf_object(page.get("/Resources") or {})
        xobjects = _pdf_object(resources.get("/XObject") or {})
        count = 0
        for value in xobjects.values():
            value = _pdf_object(value)
            if value.get("/Subtype") == "/Image":
                count += 1
        return count

    @staticmethod
    def _page_content_bytes(page: Any) -> bytes:
        contents = page.get_contents()
        if contents is None:
            return b""
        return bytes(contents.get_data())

    @staticmethod
    def _reject_unsafe_pdf(reader: PdfReader) -> None:
        if reader.is_encrypted:
            raise DocumentSecurityError("encrypted PDF content is not accepted")
        root = _pdf_object(reader.trailer.get("/Root") or {})
        names = _pdf_object(root.get("/Names") or {})
        if root.get("/OpenAction") is not None or root.get("/AA") is not None:
            raise DocumentSecurityError("active PDF actions are not accepted")
        if names.get("/JavaScript") is not None:
            raise DocumentSecurityError("PDF scripts are not accepted")
        if names.get("/EmbeddedFiles") is not None:
            raise DocumentSecurityError("embedded PDF binaries are not accepted")
        for page in reader.pages:
            if page.get("/AA") is not None:
                raise DocumentSecurityError("active PDF page actions are not accepted")
            annotations = page.get("/Annots") or ()
            for annotation in annotations:
                annotation = _pdf_object(annotation)
                if annotation.get("/Subtype") in {"/FileAttachment", "/RichMedia"}:
                    raise DocumentSecurityError("embedded PDF binaries are not accepted")
                action = annotation.get("/A")
                if action is not None:
                    action = _pdf_object(action)
                    if action.get("/S") in {"/JavaScript", "/Launch"}:
                        raise DocumentSecurityError("active PDF actions are not accepted")

    def _docx(self, path: Path) -> ParsedDocument:
        document = DocxDocument(str(path))
        blocks: list[tuple[str, int, str]] = []
        current_title = "Document"
        current_level = 1
        current_lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            heading = re.match(r"Heading\s+(\d+)", style, re.IGNORECASE)
            if heading:
                if current_lines:
                    blocks.append((current_title, current_level, "\n".join(current_lines)))
                current_title = text
                current_level = int(heading.group(1))
                current_lines = []
            else:
                current_lines.append(text)
        if current_lines or not blocks:
            blocks.append((current_title, current_level, "\n".join(current_lines)))
        sections = []
        cursor = 1
        for ordinal, (title, level, body) in enumerate(blocks, 1):
            line_count = max(1, len(body.splitlines()))
            sections.append(
                {
                    "ordinal": ordinal,
                    "level": level,
                    "title": title,
                    "content": body,
                    "start_line": cursor,
                    "end_line": cursor + line_count - 1,
                    "source_locator": f"{path}#section={ordinal}",
                }
            )
            cursor += line_count
        tables = []
        for ordinal, table in enumerate(document.tables, 1):
            matrix = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append(
                {
                    "ordinal": ordinal,
                    "title": f"Table {ordinal}",
                    "caption": "",
                    "rows": matrix,
                    "source_locator": f"{path}#table={ordinal}",
                    "extraction_confidence": 1.0,
                    "metadata": {"parser": "python-docx"},
                }
            )
        content = "\n\n".join(f"{item['title']}\n{item['content']}" for item in sections)
        figures = [
            {
                "ordinal": index,
                "caption": "",
                "source_locator": f"{path}#figure={index}",
                "artifact_ref": None,
                "metadata": {"parser": "python-docx", "inline_shape": True},
            }
            for index, _shape in enumerate(document.inline_shapes, 1)
        ]
        return ParsedDocument(
            content=content,
            source_type="docx",
            source_uri=str(path),
            sections=sections,
            tables=tables,
            figures=figures,
            citations=self._citations(content, str(path)),
            metadata={"layout_parser": "python-docx"},
        )

    def _html(self, path: Path) -> ParsedDocument:
        parser = _TextHTMLParser()
        parser.feed(path.read_text(encoding="utf-8"))
        content = "\n".join(parser.parts)
        return ParsedDocument(
            content=content,
            source_type="html",
            source_uri=str(path),
            sections=self._line_sections(content, str(path)),
            citations=self._citations(content, str(path)),
        )

    @staticmethod
    def _line_sections(content: str, source_uri: str) -> list[dict[str, Any]]:
        lines = content.splitlines()
        headings = []
        for index, line in enumerate(lines, start=1):
            match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
            if match:
                headings.append((index, len(match.group(1)), match.group(2).strip()))
        if not headings:
            headings = [(1, 1, "Document")]
        sections = []
        for ordinal, (start, level, title) in enumerate(headings, start=1):
            end = headings[ordinal][0] - 1 if ordinal < len(headings) else len(lines)
            body_start = start + 1 if lines and lines[start - 1].startswith("#") else start
            body = "\n".join(lines[body_start - 1 : end]).strip()
            sections.append(
                {
                    "ordinal": ordinal,
                    "level": level,
                    "title": title,
                    "content": body,
                    "start_line": start,
                    "end_line": max(start, end),
                    "source_locator": f"{source_uri}#L{start}-L{max(start, end)}",
                }
            )
        return sections

    @staticmethod
    def _markdown_tables(content: str, source_uri: str) -> list[dict[str, Any]]:
        lines = content.splitlines()
        tables = []
        index = 0
        while index + 1 < len(lines):
            header = lines[index].strip()
            divider = lines[index + 1].strip()
            if "|" not in header or not re.match(r"^\|?\s*:?-+", divider):
                index += 1
                continue
            rows = [StructuredDocumentAdapter._pipe_row(header)]
            cursor = index + 2
            while cursor < len(lines) and "|" in lines[cursor]:
                rows.append(StructuredDocumentAdapter._pipe_row(lines[cursor]))
                cursor += 1
            ordinal = len(tables) + 1
            tables.append(
                {
                    "ordinal": ordinal,
                    "title": f"Table {ordinal}",
                    "caption": "",
                    "rows": rows,
                    "source_locator": f"{source_uri}#L{index + 1}-L{cursor}",
                    "extraction_confidence": 0.98,
                    "metadata": {"parser": "markdown-table-v1"},
                }
            )
            index = cursor
        return tables

    @staticmethod
    def _pipe_row(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip().strip("|").split("|")]

    @staticmethod
    def _citations(content: str, source_uri: str) -> list[dict[str, Any]]:
        patterns = re.compile(r"(?:\[(\d{1,4})\]|(10\.\d{4,9}/[-._;()/:A-Z0-9]+))", re.I)
        results = []
        for ordinal, match in enumerate(patterns.finditer(content), 1):
            start = max(0, match.start() - 140)
            end = min(len(content), match.end() + 140)
            results.append(
                {
                    "ordinal": ordinal,
                    "marker": match.group(0),
                    "context": content[start:end].strip(),
                    "target": match.group(2),
                    "source_locator": f"{source_uri}#char={match.start()}-{match.end()}",
                    "metadata": {"kind": "doi" if match.group(2) else "numeric"},
                }
            )
        return results[:2000]


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)
